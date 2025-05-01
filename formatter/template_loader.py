from docx import Document
import os

class TemplateLoader:
    def __init__(self, template_path):
        self.template_path = template_path
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Template file not found: {template_path}")
        self.document = Document(template_path)

    def insert_data(self, data):
        """
        Insert parsed data into the template.
        :param data: Dictionary containing the parsed data
        """
        try:
            # Load the template
            doc = self.document

            # Iterate through paragraphs to find and replace placeholders
            for paragraph in doc.paragraphs:
                for key, value in data.items():
                    if f"{{{key}}}" in paragraph.text:
                        paragraph.text = paragraph.text.replace(f"{{{key}}}", str(value))

            return doc
        except Exception as e:
            print(f"Error inserting data into template: {e}")
            return None

    def save_document(self, output_path, doc=None):
        """
        Save the modified document
        :param output_path: Path where the document should be saved
        :param doc: Optional document object to save (if not provided, uses self.document)
        """
        try:
            document_to_save = doc if doc else self.document
            document_to_save.save(output_path)
            return True
        except Exception as e:
            print(f"Error saving document: {e}")
            return False